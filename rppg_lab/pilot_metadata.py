"""Read-only pilot metadata with explicit mappings and lossless raw tables.

No header guessing, date guessing, row-order matching, or inferred BP labels.
Google authorization is supplied by the caller, never stored in artifacts.
"""
from __future__ import annotations

import csv
from datetime import datetime, date, time
import json
from pathlib import Path
import re
import urllib.request

from .artifacts import write_json

FIELDS = ("subject_id", "session_id", "recording_id", "video_path", "recording_start",
          "recording_end", "bp_start", "bp_end", "bp_relative_start_sec", "bp_relative_end_sec",
          "sbp_mmHg", "dbp_mmHg", "hr_bpm", "camera", "fps", "resolution", "lighting",
          "movement", "orientation", "conditions", "suprasternal_face_cm", "suprasternal_finger_cm",
          "event_time_source", "manual_exclude", "hand_side", "cuff_arm", "cuff_device")
NUMERIC = {"sbp_mmHg", "dbp_mmHg", "hr_bpm", "fps", "bp_relative_start_sec", "bp_relative_end_sec",
           "suprasternal_face_cm", "suprasternal_finger_cm"}


def read_tables(path, header_rows=None, csv_sheet_name=None):
    """header_rows maps exact sheet names to zero-based header indices."""
    path = Path(path)
    header_rows = header_rows or {}
    if path.suffix.lower() == ".csv":
        name = csv_sheet_name or path.stem
        with path.open(encoding="utf-8-sig", newline="") as stream:
            matrices = {name: list(csv.reader(stream))}
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            matrices = {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in book}
        finally:
            book.close()
    else:
        raise ValueError("Metadata must be .csv or .xlsx")
    result = {}
    for name, matrix in matrices.items():
        index = header_rows.get(name, 0)
        if type(index) is not int or not 0 <= index < len(matrix):
            raise ValueError(f"Invalid explicit header row for {name}")
        # Preserve preambles, blanks, duplicate headers and every cell in raw.
        raw = [[v.isoformat() if isinstance(v, (datetime,date,time)) else v for v in row] for row in matrix]
        headers = ["" if v is None else str(v) for v in matrix[index]]
        duplicates = sorted({h for h in headers if headers.count(h) > 1})
        records = []
        if not duplicates:
            for row_index, values in enumerate(raw[index + 1:], index + 2):
                if any(v not in (None, "") for v in values):
                    records.append(dict(source_row=row_index, values=dict(zip(headers, values))))
        result[name] = dict(headers=headers, header_row=index, duplicate_headers=duplicates,
                            raw_matrix=raw, rows=records)
    return result


def fetch_google_workbook(spreadsheet_id, destination, token=None):
    """All tabs via public XLSX, or Sheets API with an authorized bearer token.

    The token needs read access to the workbook; it is never written or printed.
    """
    if not re.fullmatch(r"[A-Za-z0-9_-]+", spreadsheet_id):
        raise ValueError("Invalid spreadsheet ID")
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    if token:
        base = f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}"
        def get(url):
            request = urllib.request.Request(url, headers={"Authorization": "Bearer " + token})
            with urllib.request.urlopen(request, timeout=30) as response:
                return json.load(response)
        workbook = get(base + "?fields=sheets.properties")
        from urllib.parse import quote
        tables = {}
        for sheet in workbook["sheets"]:
            prop = sheet["properties"]
            name = prop["title"]
            area = "'" + name.replace("'", "''") + "'"
            data = get(base + "/values/" + quote(area, safe="") + "?valueRenderOption=FORMATTED_VALUE")
            tables[name] = dict(properties=prop, values=data.get("values", []))
        write_json(destination / "google_raw_tables.json", tables)
        # XLSX preserves exact worksheet names and headers for the common importer.
        from openpyxl import Workbook
        book = Workbook()
        book.remove(book.active)
        for name, table in tables.items():
            sheet = book.create_sheet(name)
            for row in table["values"]:
                sheet.append(row)
        book.save(destination / "metadata.xlsx")
    else:
        url = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=xlsx"
        with urllib.request.urlopen(url, timeout=30) as response:
            data = response.read()
        if not data.startswith(b"PK"):
            raise ValueError("Google export did not return an XLSX workbook")
        (destination / "metadata.xlsx").write_bytes(data)
    return destination / "metadata.xlsx"


def strict_timestamp(value, spec):
    if value in (None, ""):
        return None
    fmt = spec.get("format")
    if fmt == "iso8601":
        text = str(value)
        if not re.match(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}", text):
            raise ValueError("Require full ISO datetime with T and explicit timezone")
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            raise ValueError("Timezone required for ISO datetime")
    elif fmt and fmt.startswith("strptime:"):
        parsed = datetime.strptime(str(value), fmt[len("strptime:"):])
        if parsed.tzinfo is None:
            from datetime import timezone, timedelta
            offset = spec.get("utc_offset_minutes")
            if not isinstance(offset, (int, float)):
                raise ValueError("Explicit UTC offset required for naive source datetime")
            parsed = parsed.replace(tzinfo=timezone(timedelta(minutes=offset)))
    else:
        raise ValueError("Unknown timestamp format/calendar; configure it explicitly")
    return parsed.isoformat()


def normalize_tables(tables, mapping):
    normalized, issues = [], []
    for name, table in tables.items():
        cfg = mapping.get("sheets", {}).get(name)
        if cfg is None:
            issues.append(dict(sheet=name, reason="unmapped_worksheet"))
            continue
        if table["duplicate_headers"]:
            issues.append(dict(sheet=name, reason="duplicate_headers", headers=table["duplicate_headers"]))
            continue
        for item in table["rows"]:
            raw = item["values"]
            if cfg.get("required_id_column") and not raw.get(cfg["required_id_column"]):
                continue  # Empty formula/template rows remain in raw_matrix.
            record = dict(source_sheet=name, source_row=item["source_row"], raw=raw,
                          table_kind=cfg.get("kind", "unspecified"), normalized={})
            for key, spec in cfg.get("columns", {}).items():
                if key not in FIELDS:
                    raise ValueError(f"Unknown canonical metadata field: {key}")
                spec = {"column": spec} if isinstance(spec, str) else spec
                column = spec["column"]
                if column not in table["headers"]:
                    issues.append(dict(sheet=name, source_row=item["source_row"], field=key,
                                       reason="mapped_source_column_missing", column=column))
                    continue
                value = raw.get(column)
                try:
                    if value is None or str(value).strip() in ("", "NA", "N/A", "null"):
                        value = None
                    elif key in NUMERIC:
                        import math
                        value = float(value) * float(spec.get("scale", 1))
                        if not math.isfinite(value):
                            raise ValueError("Nonfinite numeric metadata")
                        if key not in ("bp_relative_start_sec", "bp_relative_end_sec") and value <= 0:
                            raise ValueError("Nonpositive physiological/acquisition value")
                    elif key in ("recording_start", "recording_end", "bp_start", "bp_end"):
                        value = strict_timestamp(value, spec)
                    elif key == "video_path":
                        value = str(value).strip().replace("\\", "/")
                    else:
                        value = str(value).strip()
                    record["normalized"][key] = value
                except (ValueError, TypeError) as exc:
                    record["normalized"][key] = None
                    issues.append(dict(sheet=name, source_row=item["source_row"], field=key,
                                       reason="invalid_value", detail=str(exc)))
            n = record["normalized"]
            if n.get("sbp_mmHg") is not None and n.get("dbp_mmHg") is not None and n["sbp_mmHg"] <= n["dbp_mmHg"]:
                n["sbp_mmHg"] = n["dbp_mmHg"] = None
                issues.append(dict(sheet=name, source_row=item["source_row"], reason="invalid_bp_order"))
            record["missing_fields"] = [k for k in FIELDS if n.get(k) is None]
            normalized.append(record)
            # Wide cuff columns become distinct measurement observations, never
            # duplicate session-matching candidates or continuous video labels.
            for index, columns in enumerate(cfg.get("measurements", []), 1):
                sub_config = {"sheets": {name: {"kind": "bp_reference", "columns": columns}}}
                sub_tables = {name: {**table, "rows": [item]}}
                children, child_issues = normalize_tables(sub_tables, sub_config)
                issues.extend(child_issues)
                for child in children:
                    child["measurement_index"] = index
                    child["normalized"] = {**{k:n.get(k) for k in ("subject_id", "session_id", "recording_id", "event_time_source", "manual_exclude", "cuff_device")},
                                           **child["normalized"]}
                    normalized.append(child)
    return normalized, issues


def match_video(video, records):
    """Strict filename convention, exact path/session matching; never row order.

    Subject rows can enrich a recording but cannot stand in for session rows.
    """
    path = Path(video)
    match = re.fullmatch(r"(S\d{3})_(V\d{2})", path.stem, re.IGNORECASE)
    identity = dict(subject_id=match[1].upper() if match else None,
                    recording_id=path.stem.upper() if match else None,
                    session_id=match[2].upper() if match else None,
                    identity_source="filename_regex_Sddd_Vdd" if match else "unresolved")
    exact = [r for r in records if r["table_kind"] == "session" and r["normalized"].get("video_path") and
             Path(r["normalized"]["video_path"]).name.casefold() == path.name.casefold()]
    candidates = exact or [r for r in records if r["table_kind"] == "session" and
                           r["normalized"].get("recording_id") == identity["recording_id"] and identity["recording_id"]]
    subject = [r for r in records if r["table_kind"] == "subject" and
               r["normalized"].get("subject_id") == identity["subject_id"] and identity["subject_id"]]
    status = "ambiguous" if len(candidates) > 1 or len(subject) > 1 else "matched_session" if candidates else "subject_only" if subject else "unmatched"
    selected = candidates[0] if len(candidates) == 1 else None
    if selected:
        for key in ("subject_id", "session_id", "recording_id"):
            mapped = selected["normalized"].get(key)
            if mapped:
                if key == "subject_id" and identity[key] and identity[key] != mapped:
                    status = "ambiguous"
                else:
                    identity[key] = mapped
    return dict(**identity, match_status=status, matched_row=selected,
                subject_row=subject[0] if len(subject) == 1 else None)


def align_references(identity, records, duration_sec):
    """One row per cuff measurement. No propagation to signal windows.

    Time spans require measured cuff start/end plus a documented association
    protocol and synchronization uncertainty in the mapping, checked by caller.
    """
    rows = []
    for record in records:
        n = record["normalized"]
        if n.get("subject_id") != identity["subject_id"] or record["table_kind"] == "subject":
            continue
        rid, sid = n.get("recording_id"), n.get("session_id")
        if rid and rid != identity["recording_id"] or sid and sid != identity["session_id"]:
            continue
        if not ("sbp_mmHg" in n or "dbp_mmHg" in n):
            continue
        row = dict(subject_id=identity["subject_id"], recording_id=identity["recording_id"],
                   source_sheet=record["source_sheet"], source_row=record["source_row"],
                   observed_cuff_sbp_mmHg=n.get("sbp_mmHg"), observed_cuff_dbp_mmHg=n.get("dbp_mmHg"),
                   hr_reference_bpm=n.get("hr_bpm"), alignment="unavailable", start_sec=None, end_sec=None,
                   reason="missing_bp", prediction_kind="observed_cuff_only")
        row["measurement_index"] = record.get("measurement_index",1)
        row["logged_start_sec"] = n.get("bp_relative_start_sec")
        row["logged_end_sec"] = n.get("bp_relative_end_sec")
        row["event_time_source"] = n.get("event_time_source")
        if n.get("sbp_mmHg") is not None and n.get("dbp_mmHg") is not None:
            row.update(alignment="session_only" if rid or sid else "subject_only",
                       reason="insufficient_timestamps_or_protocol")
            a,b = n.get("bp_relative_start_sec"), n.get("bp_relative_end_sec")
            start = identity.get("matched_row")
            video_start = start["normalized"].get("recording_start") if start else None
            if a is None and video_start and n.get("bp_start") and n.get("bp_end"):
                origin = datetime.fromisoformat(video_start)
                a = (datetime.fromisoformat(n["bp_start"]) - origin).total_seconds()
                b = (datetime.fromisoformat(n["bp_end"]) - origin).total_seconds()
            if a is not None and b is not None:
                if 0 <= a < b <= duration_sec:
                    verified = n.get("event_time_source") in ("Timer", "Video_review")
                    row.update(alignment="timed_cuff_interval" if verified else "session_only",
                               start_sec=a if verified else None, end_sec=b if verified else None,
                               reason="cuff_measurement_interval_only_not_continuous_ground_truth" if verified else "legacy_event_times_need_video_confirmation")
                else:
                    row["reason"] = "cuff_interval_outside_recording_or_invalid"
        if identity["match_status"] == "ambiguous":
            row.update(alignment="unavailable", reason="ambiguous_video_match", start_sec=None, end_sec=None)
        rows.append(row)
    return rows
